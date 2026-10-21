import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from scripts.build_dashboard_snapshot import build_snapshot


class DashboardSnapshotTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.tempdir.name) / "Keuangan.xlsx"
        wb = Workbook()
        dashboard = wb.active
        dashboard.title = "Dashboard"
        dashboard.cell(25, 1, "Dompet")
        dashboard.cell(25, 2, "Saldo Sistem")
        balances = {"CV HB(101)": 1000, "TX SBY(216)": 2000, "TX BALI(087)": 3000}
        for row, (wallet, balance) in enumerate(balances.items(), start=26):
            dashboard.cell(row, 1, wallet)
            dashboard.cell(row, 2, balance)

        for wallet in balances:
            ws = wb.create_sheet(wallet)
            headers = ["No", "Waktu", "Waktu/Tanggal", "Jumlah", "Project", "Keterangan", "Oleh", "Source", "MessageID"]
            for col, value in enumerate(headers, start=1):
                ws.cell(8, col, value)
                ws.cell(8, col + 9, value)

        holla = wb["CV HB(101)"]
        self._write_transaction(holla, 9, 1, 1, "2026-09-01", 5000, "Project A", "DP", "income-1")
        self._write_transaction(holla, 9, 10, 1, "2026-09-02", 1200, "Project A", "Bahan", "expense-1")
        self._write_transaction(holla, 10, 10, 2, "2026-09-02", 700, "Saldo Umum", "Hutang ke dompet TX SBY", "debt-1|UTANG")
        sby = wb["TX SBY(216)"]
        self._write_transaction(sby, 9, 1, 1, "2026-09-02", 700, "Saldo Umum", "Terima hutang", "debt-1|TERIMA")

        debts = wb.create_sheet("HUTANG")
        debts.append(["No", "Tanggal", "Nominal", "Keterangan", "Yang Hutang", "Yang Dihutangi", "Status", "Tgl Lunas", "MessageID"])
        debts.append([1, "2026-09-02", 700, "Pinjaman", "CV HB(101)", "TX SBY(216)", "OPEN", None, "debt-1|HUTANG"])
        debts.append([2, "2026-09-02", 300, "Lunas", "CV HB(101)", "TX SBY(216)", "PAID", "2026-09-02", "debt-2|HUTANG"])
        wb.save(self.workbook_path)

    @staticmethod
    def _write_transaction(ws, row, start_col, number, date, amount, project, notes, message_id):
        values = [number, "10:00", date, amount, project, notes, "Owner", "WhatsApp", message_id]
        for offset, value in enumerate(values):
            ws.cell(row, start_col + offset, value)

    def tearDown(self):
        self.tempdir.cleanup()

    def test_builds_exact_read_only_snapshot(self):
        snapshot = build_snapshot(self.workbook_path, source_modified_at="2026-10-20T16:00:00+07:00")
        metrics = snapshot["overview"]["metrics"]
        self.assertEqual(metrics["cash_balance"], 6000)
        self.assertEqual(metrics["operating_inflow"], 5000)
        self.assertEqual(metrics["operating_outflow"], 1200)
        self.assertEqual(metrics["net_operating_cashflow"], 3800)
        self.assertEqual(metrics["open_internal_debt"], 700)
        self.assertEqual(snapshot["meta"]["data_through"], "2026-09-02")
        self.assertEqual(snapshot["meta"]["freshness"], "stale")
        self.assertEqual(len(snapshot["transactions"]), 4)
        self.assertEqual(sum(1 for row in snapshot["transactions"] if row["internal_movement"]), 2)
        self.assertEqual(snapshot["wallets"][0]["balance"], 1000)
        self.assertEqual(snapshot["debts"][0]["status"], "OPEN")
        self.assertEqual(snapshot["project_summaries"][0], {"name": "Project A", "net_cashflow": 3800})
        self.assertEqual(snapshot["cash_trend"][-1]["balance"], 6000)

    def test_snapshot_is_json_serializable_and_contains_no_source_path(self):
        snapshot = build_snapshot(self.workbook_path, source_modified_at="2026-10-20T16:00:00+07:00")
        serialized = json.dumps(snapshot)
        self.assertNotIn(str(self.workbook_path), serialized)
        self.assertEqual(snapshot["version"], "dashboard.snapshot.v1")


if __name__ == "__main__":
    unittest.main()
